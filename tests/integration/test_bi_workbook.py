import csv
from pathlib import Path

from openpyxl import load_workbook

from fx_intelligence.warehouse.workbook import export_workbook


def test_powerbi_workbook_uses_named_tables_and_typed_dates(tmp_path: Path):
    data = tmp_path / "data"
    data.mkdir()
    with (data / "dim_date.csv").open("w", newline="", encoding="utf-8-sig") as handle:
        writer = csv.writer(handle)
        writer.writerow(["date_key", "year"])
        writer.writerow(["2026-01-02", "2026"])
    path = export_workbook(data, tmp_path / "model.xlsx")
    workbook = load_workbook(path, read_only=False, data_only=True)
    sheet = workbook["dim_date"]
    assert len(sheet.tables) == 1
    assert sheet["A2"].value.year == 2026
    assert sheet["B2"].value == 2026
