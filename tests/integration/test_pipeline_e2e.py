"""Reproduce the complete offline FX-to-Power-BI sample path."""

import csv
import json
from dataclasses import replace
from datetime import date
from decimal import Decimal
from pathlib import Path

import duckdb
import pytest
from openpyxl import load_workbook

from fx_intelligence.analytics.macro import available_pairs
from fx_intelligence.ingestion.normalize import normalize_payload
from fx_intelligence.registry import load_registry
from fx_intelligence.transformation.alignment import align_macro
from fx_intelligence.transformation.fx import cross_rate, invert_fx
from fx_intelligence.validation.validate import validate_fx, validate_macro
from fx_intelligence.warehouse.build import build_warehouse
from fx_intelligence.warehouse.export import export_powerbi
from fx_intelligence.warehouse.workbook import export_workbook

FIXTURE = Path("tests/fixtures/sample_pipeline.json")


def _read_csv(path: Path) -> list[dict[str, str]]:
    with path.open(newline="", encoding="utf-8-sig") as handle:
        return list(csv.DictReader(handle))


def _build_sample(root: Path):
    payloads = json.loads(FIXTURE.read_text(encoding="utf-8"))
    registry = {item["id"]: item for item in load_registry(Path("config/sources/registry.yaml"))}
    fx_rows = []
    macro_rows = []
    for source_id, envelope in payloads.items():
        fx, macro = normalize_payload(envelope, registry[source_id])
        valid_fx, fx_issues = validate_fx(fx)
        valid_macro, macro_issues = validate_macro(macro)
        assert not fx_issues and not macro_issues, source_id
        fx_rows.extend(valid_fx)
        macro_rows.extend(valid_macro)
    db_path = root / "fx.duckdb"
    counts = build_warehouse(db_path, fx_rows, macro_rows, sql_root=Path("sql"))
    out = root / "powerbi"
    exports = export_powerbi(db_path, out)
    workbook = export_workbook(out, out / "fx_model.xlsx")
    return fx_rows, macro_rows, db_path, out, workbook, counts, exports


def test_offline_pipeline_reconciles_quotes_lineage_and_bi_export(tmp_path: Path) -> None:
    fx, macro, db, out, workbook, counts, exports = _build_sample(tmp_path)
    assert counts == {"fx": 9, "macro": 2}
    assert exports["fact_fx_rates"] == 9
    assert exports["fact_macro"] == 2
    assert len(_read_csv(out / "fact_fx_metrics.csv")) == 9
    with duckdb.connect(str(db), read_only=True) as conn:
        assert conn.execute("SELECT COUNT(*) FROM fact_fx_rates").fetchone()[0] == 9
        assert conn.execute("SELECT COUNT(*) FROM fact_macro").fetchone()[0] == 2
        assert conn.execute("SELECT COUNT(*) FROM dim_fx_pair").fetchone()[0] == 2
        assert conn.execute("SELECT COUNT(*) FROM v_fx_latest").fetchone()[0] == 3
    facts = _read_csv(out / "fact_fx_rates.csv")
    assert (
        len(
            {(r["source_id"], r["series_id"], r["observation_date"], r["rate_type"]) for r in facts}
        )
        == 9
    )
    interbank = [row for row in facts if row["source_id"] == "bot_usdthb_interbank"]
    assert [row["rate"] for row in interbank] == ["40.0000000000", "44.0000000000", "48.0000000000"]
    metrics = _read_csv(out / "fact_fx_metrics.csv")
    second = next(
        row
        for row in metrics
        if row["source_id"] == "bot_usdthb_interbank" and row["observation_date"] == "2025-01-03"
    )
    assert Decimal(second["indexed_rate"]) == Decimal("110")
    assert Decimal(second["simple_return"]) == Decimal("0.1")
    usd_interbank = next(
        row
        for row in fx
        if row.source_id == "bot_usdthb_interbank" and row.observation_date == date(2025, 1, 3)
    )
    usd_bank = next(
        row
        for row in fx
        if row.source_id == "bot_usdthb_bank_average" and row.observation_date == date(2025, 1, 3)
    )
    eur_bank = next(
        row
        for row in fx
        if row.source_id == "bot_eurthb_bank_average" and row.observation_date == date(2025, 1, 3)
    )
    assert invert_fx(usd_interbank).rate * usd_interbank.rate == pytest.approx(1)
    with pytest.raises(ValueError, match="same rate type"):
        cross_rate(usd_interbank, eur_bank)
    with pytest.raises(ValueError, match="same source"):
        cross_rate(usd_bank, eur_bank)
    assert len(macro) == 2
    book = load_workbook(workbook, read_only=True, data_only=True)
    assert "fact_fx_rates" in book.sheetnames
    assert book["fact_fx_rates"].max_row == 10
    book.close()


def test_no_macro_future_leakage_when_release_date_is_unknown_or_later(tmp_path: Path) -> None:
    _, macro, *_ = _build_sample(tmp_path)
    assert all(row.release_date is None for row in macro)
    returns = [(date(2025, 1, 3), 0.1), (date(2025, 1, 6), 0.09)]
    assert available_pairs(returns, macro) == []
    released = replace(macro[0], release_date=date(2025, 1, 10))
    aligned = align_macro([released], [date(2025, 1, 31), date(2025, 2, 1)])
    assert aligned == [None, released]


def test_two_clean_sample_builds_export_identical_observations(tmp_path: Path) -> None:
    first = _build_sample(tmp_path / "first")[3]
    second = _build_sample(tmp_path / "second")[3]
    files = sorted(path.name for path in first.glob("*.csv") if path.name != "refresh_metadata.csv")
    assert files
    assert all((first / name).read_bytes() == (second / name).read_bytes() for name in files)
