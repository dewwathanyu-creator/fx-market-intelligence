"""Package generated CSV extracts as one typed Excel source for Power BI Desktop."""

import csv
from datetime import date
from pathlib import Path

from openpyxl import Workbook
from openpyxl.styles import Font
from openpyxl.utils import get_column_letter
from openpyxl.worksheet.table import Table, TableStyleInfo

DATE_COLUMNS = {"date_key", "observation_date", "vintage_date", "release_date"}
INTEGER_COLUMNS = {"year", "month", "day", "fx_row_count", "macro_row_count"}
NUMERIC_COLUMNS = {
    "rate",
    "value",
    "indexed_rate",
    "simple_return",
    "log_return",
    "drawdown",
    "rolling_return_5",
    "rolling_return_20",
    "rolling_return_60",
    "volatility_20",
    "volatility_60",
    "volatility_120",
}


def _typed(value: str, column: str) -> object:
    if value == "":
        return None
    if column in DATE_COLUMNS:
        return date.fromisoformat(value)
    if column in INTEGER_COLUMNS:
        return int(value)
    if column in NUMERIC_COLUMNS:
        return float(value)
    return value


def export_workbook(csv_dir: Path, output: Path) -> Path:
    """Create one Excel table per CSV, preserving date and numeric column types."""
    files = sorted(csv_dir.glob("*.csv"))
    if not files:
        raise ValueError("no BI CSV exports found")
    workbook = Workbook()
    workbook.remove(workbook.active)
    for path in files:
        with path.open(newline="", encoding="utf-8-sig") as handle:
            reader = csv.DictReader(handle)
            headers = reader.fieldnames
            if not headers:
                raise ValueError(f"empty CSV headers: {path.name}")
            sheet = workbook.create_sheet(path.stem)
            sheet.append(headers)
            for record in reader:
                sheet.append([_typed(record[column], column) for column in headers])
        sheet.freeze_panes = "A2"
        sheet.auto_filter.ref = sheet.dimensions
        for cell in sheet[1]:
            cell.font = Font(bold=True)
        for row in sheet.iter_rows(min_row=2):
            for cell, column in zip(row, headers, strict=True):
                if column in DATE_COLUMNS and cell.value is not None:
                    cell.number_format = "yyyy-mm-dd"
        if sheet.max_row > 1:
            table = Table(
                displayName=f"tbl_{path.stem}",
                ref=f"A1:{get_column_letter(len(headers))}{sheet.max_row}",
            )
            table.tableStyleInfo = TableStyleInfo(name="TableStyleMedium2", showRowStripes=True)
            sheet.add_table(table)
    output.parent.mkdir(parents=True, exist_ok=True)
    workbook.save(output)
    return output
